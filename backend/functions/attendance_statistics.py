"""
Attendance statistics for the exported Google Sheets.

This module builds the "Statistics" worksheet that accompanies the raw
attendance grids written by ``attendance_sheets.py`` (Sunday School) and
``servants_attendance_sheets.py`` (Servants Meeting).

Stats are computed **per person**, one row each, grouped into sections
(one section per grade for Sunday School, a single section for the
Servants Meeting). Every row answers three questions:

  1. When they were marked, how often were they there?   -> Attendance %
  2. How often were they marked at all?                  -> Marked % / Unaccounted %
  3. Is their pattern steady or erratic?                  -> Consistency % / Pattern

Attendance is stored in Firestore as ``"MM-DD-YYYY": bool`` fields on the
person's document. ``True`` = Present, ``False`` = Absent, and a *missing*
field means nobody recorded anything that week - that is what "unaccounted
for" means throughout this module. The distinction matters: an absent
student was seen to be missing, an unaccounted student was never checked.
"""

import datetime


# === WORKSHEET NAMING ===
STATS_SHEET_TITLE = "Statistics"

# Google Sheets caps a single batchUpdate payload, and colouring cells
# individually generates a lot of requests. Send them in chunks.
REQUEST_CHUNK_SIZE = 1000


# === COLOR PALETTE ===
# Soft tints rather than the pure green/red used on the attendance grids -
# these cells contain text that still has to be readable on top of the fill.
def _rgb(r, g, b):
    return {"red": r / 255.0, "green": g / 255.0, "blue": b / 255.0}


GREEN_STRONG = _rgb(146, 208, 148)
GREEN_LIGHT = _rgb(200, 230, 201)
YELLOW = _rgb(255, 241, 176)
ORANGE = _rgb(252, 214, 176)
RED = _rgb(244, 187, 184)
GREY = _rgb(235, 235, 235)

WHITE = _rgb(255, 255, 255)
HEADER_BG = _rgb(20, 126, 251)        # Taksis accent blue
SECTION_BG = _rgb(38, 38, 38)         # dark band that opens each grade
SUMMARY_BG = _rgb(222, 235, 250)      # tinted band for the grade average
TITLE_BG = _rgb(15, 100, 210)

WHITE_TEXT = {"foregroundColor": WHITE}
DARK_TEXT = {"foregroundColor": _rgb(0, 0, 0)}


# === COLUMN LAYOUT ===
# (header, stats key, colour mode, pixel width)
#   "none"       - no fill
#   "pct"        - green when high, red when low
#   "pct_invert" - green when low, red when high (unaccounted %)
#   "label"      - fill looked up from LABEL_COLORS
COLUMNS = [
    ("Name",             "name",             "none",       190),
    ("Weeks Tracked",    "weeks",            "none",        95),
    ("Present",          "present",          "none",        75),
    ("Absent",           "absent",           "none",        75),
    ("Unaccounted",      "unaccounted",      "none",       100),
    ("Attendance %",     "attendance_pct",   "pct",        105),
    ("Marked %",         "marked_pct",       "pct",         90),
    ("Unaccounted %",    "unaccounted_pct",  "pct_invert", 110),
    ("Consistency %",    "consistency_pct",  "pct",        105),
    ("Pattern",          "pattern",          "label",      150),
    ("Current Streak",   "current_streak",   "none",       110),
    ("Longest Present",  "longest_present",  "none",       110),
    ("Last 4 Marked %",  "recent_pct",       "pct",        115),
    ("Last Present",     "last_present",     "none",       105),
    ("Status",           "status",           "label",      100),
]

PERCENT_COLUMN_INDEXES = [
    i for i, col in enumerate(COLUMNS) if col[2] in ("pct", "pct_invert")
]

# Column widths for the printed .xlsx, in Excel character units.
#
# Deliberately much narrower than the on-screen pixel widths above: the full
# set has to fit across one landscape page, and scaling down to reach that is
# what makes a printout unreadable. These total ~150 units (~11in), so the
# fit-to-width scale lands around 90% rather than 65%. Headers wrap onto two
# lines, which is what lets the columns be this narrow.
PRINT_WIDTHS = {
    "name": 20,
    "weeks": 7,
    "present": 8,
    "absent": 8,
    "unaccounted": 9,
    "attendance_pct": 9,
    "marked_pct": 9,
    "unaccounted_pct": 9,
    "consistency_pct": 9,
    "pattern": 16,
    "current_streak": 10,
    "longest_present": 9,
    "recent_pct": 9,
    "last_present": 11,
    "status": 10,
}

NUM_COLUMNS = len(COLUMNS)

LABEL_COLORS = {
    # Status
    "Excellent": GREEN_STRONG,
    "Good": GREEN_LIGHT,
    "Watch": YELLOW,
    "At Risk": ORANGE,
    "Critical": RED,
    "No Data": GREY,
    # Pattern
    "Consistently Present": GREEN_STRONG,
    "Mostly Present": GREEN_LIGHT,
    "Mixed": YELLOW,
    "Erratic": ORANGE,
    "Consistently Absent": RED,
}

LEGEND = (
    "Attendance % = Present ÷ weeks actually marked   |   "
    "Marked % = weeks marked ÷ weeks tracked (the rest are unaccounted for)   |   "
    "Consistency % = how rarely the person flips between Present and Absent   |   "
    "Green = good, red = needs attention. A blank cell means there is no data to judge."
)


# === STAT COMPUTATION ===
def _percent(numerator, denominator):
    """Return a 0-100 percentage, or None when there is nothing to divide by."""
    if not denominator:
        return None
    return (numerator / float(denominator)) * 100.0


def _classify_pattern(attendance_pct, consistency_pct, recorded):
    """Describe the shape of someone's attendance in plain words.

    Consistency alone is ambiguous - somebody who never shows up is
    perfectly consistent - so the label always folds in attendance % to say
    *what* they are consistent about.
    """
    if recorded == 0:
        return "No Data"
    if consistency_pct is None:
        # A single recorded week: not enough to talk about a pattern.
        return "Mostly Present" if attendance_pct >= 50 else "Mixed"
    if consistency_pct >= 70:
        if attendance_pct >= 70:
            return "Consistently Present"
        if attendance_pct <= 30:
            return "Consistently Absent"
        return "Mixed"
    if consistency_pct < 40:
        return "Erratic"
    return "Mostly Present" if attendance_pct >= 60 else "Mixed"


def _classify_status(attendance_pct, marked_pct, recorded):
    """Overall flag for the row, driven by attendance with a data-quality guard."""
    if recorded == 0:
        return "No Data"
    # Attendance computed from a couple of marked weeks out of thirty is not
    # trustworthy enough to call "Excellent", so cap it.
    thin_data = marked_pct is not None and marked_pct < 40
    if attendance_pct >= 90:
        return "Good" if thin_data else "Excellent"
    if attendance_pct >= 75:
        return "Watch" if thin_data else "Good"
    if attendance_pct >= 60:
        return "Watch"
    if attendance_pct >= 40:
        return "At Risk"
    return "Critical"


def compute_person_stats(name, attendance, date_array):
    """Build the statistics row for one person.

    Args:
        name: display name for the row.
        attendance: dict of ``"MM-DD-YYYY" -> value``. Only ``True``/``False``
            count; anything else (missing, None, a stray string) is treated as
            unaccounted for, matching how the attendance grids read the data.
        date_array: ordered list of the ``"MM-DD-YYYY"`` week keys the sheet
            covers. This defines the denominator - weeks outside it are ignored.

    Returns:
        dict keyed by the ``COLUMNS`` stat keys.
    """
    weeks = len(date_array)

    # Walk the weeks in order, keeping only the ones somebody actually marked.
    marked_sequence = []
    present = 0
    absent = 0
    last_present = None

    for date_key in date_array:
        value = attendance.get(date_key)
        if value is True:
            present += 1
            marked_sequence.append(True)
            last_present = date_key
        elif value is False:
            absent += 1
            marked_sequence.append(False)

    recorded = present + absent
    unaccounted = weeks - recorded

    attendance_pct = _percent(present, recorded)
    marked_pct = _percent(recorded, weeks)
    unaccounted_pct = _percent(unaccounted, weeks)

    # Consistency: how rarely the status flips between one marked week and the
    # next. Needs at least two marked weeks to mean anything.
    consistency_pct = None
    if len(marked_sequence) >= 2:
        switches = sum(
            1 for i in range(1, len(marked_sequence))
            if marked_sequence[i] != marked_sequence[i - 1]
        )
        consistency_pct = 100.0 - _percent(switches, len(marked_sequence) - 1)

    # Current streak: the trailing run of identical marked weeks.
    current_streak = "—"
    if marked_sequence:
        streak_value = marked_sequence[-1]
        streak_length = 0
        for value in reversed(marked_sequence):
            if value != streak_value:
                break
            streak_length += 1
        current_streak = "%d %s" % (
            streak_length, "Present" if streak_value else "Absent"
        )

    # Longest run of consecutive Presents.
    longest_present = 0
    run = 0
    for value in marked_sequence:
        run = run + 1 if value else 0
        longest_present = max(longest_present, run)

    # Recent trend over the last four *marked* weeks.
    recent_window = marked_sequence[-4:]
    recent_pct = _percent(sum(1 for v in recent_window if v), len(recent_window))

    return {
        "name": name,
        "weeks": weeks,
        "present": present,
        "absent": absent,
        "unaccounted": unaccounted,
        "attendance_pct": attendance_pct,
        "marked_pct": marked_pct,
        "unaccounted_pct": unaccounted_pct,
        "consistency_pct": consistency_pct,
        "pattern": _classify_pattern(attendance_pct or 0, consistency_pct, recorded),
        "current_streak": current_streak,
        "longest_present": longest_present,
        "recent_pct": recent_pct,
        "last_present": last_present or "Never",
        "status": _classify_status(attendance_pct or 0, marked_pct, recorded),
    }


def summarize_rows(label, rows):
    """Aggregate already-computed person rows into one summary row.

    Percentages are pooled from the raw counts rather than averaged across
    people, so a class of thirty is not swung by one student with two marked
    weeks. Consistency has no counts to pool, so it is a plain mean over the
    people who have a value.
    """
    if not rows:
        return None

    weeks = max(row["weeks"] for row in rows)
    present = sum(row["present"] for row in rows)
    absent = sum(row["absent"] for row in rows)
    unaccounted = sum(row["unaccounted"] for row in rows)
    recorded = present + absent
    total_slots = present + absent + unaccounted

    consistencies = [
        row["consistency_pct"] for row in rows if row["consistency_pct"] is not None
    ]
    consistency_pct = (
        sum(consistencies) / len(consistencies) if consistencies else None
    )

    attendance_pct = _percent(present, recorded)
    marked_pct = _percent(recorded, total_slots)

    return {
        "name": label,
        "weeks": weeks,
        "present": present,
        "absent": absent,
        "unaccounted": unaccounted,
        "attendance_pct": attendance_pct,
        "marked_pct": marked_pct,
        "unaccounted_pct": _percent(unaccounted, total_slots),
        "consistency_pct": consistency_pct,
        "pattern": _classify_pattern(
            attendance_pct or 0, consistency_pct, recorded
        ),
        "current_streak": "%d people" % len(rows),
        "longest_present": max(row["longest_present"] for row in rows),
        "recent_pct": None,
        "last_present": "—",
        "status": _classify_status(attendance_pct or 0, marked_pct, recorded),
    }


def build_person_rows(people, date_array):
    """Compute and alphabetise stats for a ``{name: attendance_dict}`` mapping."""
    return [
        compute_person_stats(name, people.get(name) or {}, date_array)
        for name in sorted(people.keys())
    ]


# === SHEET RENDERING ===
def _value_cell(value, is_percent):
    """Wrap a python value in a Sheets ``userEnteredValue``."""
    if value is None or value == "":
        return {"userEnteredValue": {"stringValue": ""}}
    if is_percent:
        # Stored as a fraction so the PERCENT number format renders it and
        # sorting/filtering in the sheet still behaves numerically.
        return {"userEnteredValue": {"numberValue": round(value / 100.0, 4)}}
    if isinstance(value, bool):
        return {"userEnteredValue": {"stringValue": str(value)}}
    if isinstance(value, (int, float)):
        return {"userEnteredValue": {"numberValue": value}}
    return {"userEnteredValue": {"stringValue": str(value)}}


def _scale_color(pct, invert=False):
    """Map a 0-100 percentage onto the green-to-red scale."""
    if pct is None:
        return GREY
    score = (100.0 - pct) if invert else pct
    if score >= 90:
        return GREEN_STRONG
    if score >= 75:
        return GREEN_LIGHT
    if score >= 60:
        return YELLOW
    if score >= 40:
        return ORANGE
    return RED


def _cell_color(stats, column):
    """Background fill for one cell, or None to leave it unfilled."""
    _, key, mode, _width = column
    value = stats.get(key)
    if mode == "pct":
        return _scale_color(value)
    if mode == "pct_invert":
        return _scale_color(value, invert=True)
    if mode == "label":
        return LABEL_COLORS.get(value, GREY)
    return None


class _Layout:
    """Accumulates the grid of values plus the formatting that goes with it."""

    def __init__(self):
        self.values = []          # list of rows, each a list of (value, is_percent)
        self.fills = []           # (row, col, color)
        self.title_rows = []
        self.section_rows = []
        self.header_rows = []
        self.summary_rows = []
        self.subblock_rows = []
        self.spacer_rows = []

    @property
    def row_count(self):
        return len(self.values)

    def blank_row(self):
        self.values.append([("", False)] * NUM_COLUMNS)
        self.spacer_rows.append(self.row_count - 1)

    def text_row(self, text):
        row = [(text, False)] + [("", False)] * (NUM_COLUMNS - 1)
        self.values.append(row)
        return self.row_count - 1

    def title_row(self, text):
        self.title_rows.append(self.text_row(text))

    def section_row(self, text):
        self.section_rows.append(self.text_row(text))

    def subblock_row(self, text):
        self.subblock_rows.append(self.text_row(text))

    def header_row(self):
        self.values.append([(col[0], False) for col in COLUMNS])
        self.header_rows.append(self.row_count - 1)

    def stats_row(self, stats, is_summary=False):
        row = []
        row_index = self.row_count
        for col_index, column in enumerate(COLUMNS):
            _, key, mode, _width = column
            is_percent = mode in ("pct", "pct_invert")
            row.append((stats.get(key), is_percent))
            color = _cell_color(stats, column)
            if color is not None:
                self.fills.append((row_index, col_index, color))
        self.values.append(row)
        if is_summary:
            self.summary_rows.append(row_index)


def build_statistics_layout(sections, title, subtitle):
    """Lay out the whole worksheet.

    Args:
        sections: ordered list of dicts::

            {
              "title": "3rd Grade",
              "blocks": [("Students", [stats, ...]), ("Servants", [stats, ...])],
              "summary_label": "3rd Grade — Overall",   # optional
            }

        title: banner text for row 1.
        subtitle: smaller line under the banner (school year, generated-at).

    Returns:
        A ``_Layout``.
    """
    layout = _Layout()
    layout.title_row(title)
    layout.text_row(subtitle)
    layout.text_row(LEGEND)
    layout.blank_row()

    for section in sections:
        blocks = [
            (label, rows) for label, rows in section.get("blocks", []) if rows
        ]
        if not blocks:
            continue

        layout.section_row(section["title"])

        for label, rows in blocks:
            # Only name the block when a section has more than one, otherwise
            # the label is just noise above the column headers.
            if len(blocks) > 1:
                layout.subblock_row(label)
            layout.header_row()
            for stats in rows:
                layout.stats_row(stats)

        # One pooled summary across every block in the section.
        all_rows = [stats for _label, rows in blocks for stats in rows]
        summary = summarize_rows(
            section.get("summary_label") or ("%s — Overall" % section["title"]),
            all_rows,
        )
        if summary:
            layout.stats_row(summary, is_summary=True)

        layout.blank_row()

    return layout


def _layout_to_requests(sheet_id, layout, total_cols, clear_rows):
    """Turn a laid-out grid into Google Sheets API requests."""
    requests = []

    # 1. Wipe whatever the previous run left behind - the row count shrinks
    #    when people leave, and stale rows below the new data would survive.
    requests.append({
        "repeatCell": {
            "range": {
                "sheetId": sheet_id,
                "startRowIndex": 0,
                "endRowIndex": clear_rows,
                "startColumnIndex": 0,
                "endColumnIndex": total_cols,
            },
            "cell": {
                "userEnteredValue": {"stringValue": ""},
                "userEnteredFormat": {
                    "backgroundColor": WHITE,
                    "textFormat": {"bold": False, "foregroundColor": _rgb(0, 0, 0)},
                    "horizontalAlignment": "CENTER",
                },
            },
            "fields": (
                "userEnteredValue,"
                "userEnteredFormat.backgroundColor,"
                "userEnteredFormat.textFormat,"
                "userEnteredFormat.horizontalAlignment"
            ),
        }
    })

    # 2. Write every value in one shot.
    requests.append({
        "updateCells": {
            "range": {
                "sheetId": sheet_id,
                "startRowIndex": 0,
                "endRowIndex": layout.row_count,
                "startColumnIndex": 0,
                "endColumnIndex": NUM_COLUMNS,
            },
            "rows": [
                {"values": [_value_cell(value, is_pct) for value, is_pct in row]}
                for row in layout.values
            ],
            "fields": "userEnteredValue",
        }
    })

    def band(row_index, background, bold, text_format_color, alignment="LEFT"):
        return {
            "repeatCell": {
                "range": {
                    "sheetId": sheet_id,
                    "startRowIndex": row_index,
                    "endRowIndex": row_index + 1,
                    "startColumnIndex": 0,
                    "endColumnIndex": NUM_COLUMNS,
                },
                "cell": {
                    "userEnteredFormat": {
                        "backgroundColor": background,
                        "textFormat": dict(
                            {"bold": bold}, **{"foregroundColor": text_format_color}
                        ),
                        "horizontalAlignment": alignment,
                    }
                },
                "fields": (
                    "userEnteredFormat.backgroundColor,"
                    "userEnteredFormat.textFormat,"
                    "userEnteredFormat.horizontalAlignment"
                ),
            }
        }

    # 3. Banner, section bands, column headers, summary rows.
    for row_index in layout.title_rows:
        requests.append(band(row_index, TITLE_BG, True, WHITE, "LEFT"))
    for row_index in layout.section_rows:
        requests.append(band(row_index, SECTION_BG, True, WHITE, "LEFT"))
    for row_index in layout.subblock_rows:
        requests.append(band(row_index, GREY, True, _rgb(0, 0, 0), "LEFT"))
    for row_index in layout.header_rows:
        requests.append(band(row_index, HEADER_BG, True, WHITE, "CENTER"))
    for row_index in layout.summary_rows:
        requests.append(band(row_index, SUMMARY_BG, True, _rgb(0, 0, 0), "CENTER"))

    # 4. Per-cell colour coding. Applied after the bands so the summary row's
    #    own green/red cells win over its tinted band.
    for row_index, col_index, color in layout.fills:
        requests.append({
            "repeatCell": {
                "range": {
                    "sheetId": sheet_id,
                    "startRowIndex": row_index,
                    "endRowIndex": row_index + 1,
                    "startColumnIndex": col_index,
                    "endColumnIndex": col_index + 1,
                },
                "cell": {"userEnteredFormat": {"backgroundColor": color}},
                "fields": "userEnteredFormat.backgroundColor",
            }
        })

    # 5. Percent columns render as "87%" rather than 0.87.
    for col_index in PERCENT_COLUMN_INDEXES:
        requests.append({
            "repeatCell": {
                "range": {
                    "sheetId": sheet_id,
                    "startRowIndex": 0,
                    "endRowIndex": layout.row_count,
                    "startColumnIndex": col_index,
                    "endColumnIndex": col_index + 1,
                },
                "cell": {
                    "userEnteredFormat": {
                        "numberFormat": {"type": "PERCENT", "pattern": "0%"}
                    }
                },
                "fields": "userEnteredFormat.numberFormat",
            }
        })

    # 6. Names stay left-aligned and bold; everything else is centred.
    requests.append({
        "repeatCell": {
            "range": {
                "sheetId": sheet_id,
                "startRowIndex": 0,
                "endRowIndex": layout.row_count,
                "startColumnIndex": 0,
                "endColumnIndex": 1,
            },
            "cell": {
                "userEnteredFormat": {
                    "textFormat": {"bold": True},
                    "horizontalAlignment": "LEFT",
                }
            },
            "fields": (
                "userEnteredFormat.textFormat.bold,"
                "userEnteredFormat.horizontalAlignment"
            ),
        }
    })

    # 7. Column widths and a frozen name column.
    for col_index, column in enumerate(COLUMNS):
        requests.append({
            "updateDimensionProperties": {
                "range": {
                    "sheetId": sheet_id,
                    "dimension": "COLUMNS",
                    "startIndex": col_index,
                    "endIndex": col_index + 1,
                },
                "properties": {"pixelSize": column[3]},
                "fields": "pixelSize",
            }
        })

    requests.append({
        "updateSheetProperties": {
            "properties": {
                "sheetId": sheet_id,
                "gridProperties": {"frozenColumnCount": 1},
            },
            "fields": "gridProperties.frozenColumnCount",
        }
    })

    return requests


def write_statistics_worksheet(worksheet, sections, title, subtitle,
                               ensure_size=None, retry=None):
    """Render the statistics layout onto ``worksheet``.

    Args:
        worksheet: gspread worksheet for the Statistics tab.
        sections: see :func:`build_statistics_layout`.
        title / subtitle: banner text.
        ensure_size: optional ``fn(worksheet, rows, cols)`` used to grow the
            worksheet before writing (the callers already have one).
        retry: optional ``fn(callable)`` wrapper for the batch update.

    Returns:
        True when something was written, False otherwise.
    """
    layout = build_statistics_layout(sections, title, subtitle)
    # No header row means no section had any people in it - writing a
    # banner over the previous run's numbers would be worse than doing
    # nothing.
    if not layout.header_rows:
        return False

    required_rows = layout.row_count + 20
    required_cols = NUM_COLUMNS + 2

    # Grow the worksheet first. Every range below is validated against the
    # real grid, so writing past the current bounds is an API error rather
    # than an auto-expand.
    if ensure_size:
        ensure_size(worksheet, required_rows, required_cols)
    else:
        try:
            if worksheet.row_count < required_rows or worksheet.col_count < required_cols:
                worksheet.resize(
                    rows=max(worksheet.row_count, required_rows),
                    cols=max(worksheet.col_count, required_cols),
                )
        except Exception:
            pass

    # Clear the full grid so rows dropped since the last run do not linger
    # below the new data. Use the worksheet's real dimensions: a range past
    # the end of the grid is rejected by the API, and if the resize above did
    # not take effect these are the only bounds we can trust.
    try:
        clear_rows = worksheet.row_count
        total_cols = worksheet.col_count
    except Exception:
        clear_rows = required_rows
        total_cols = required_cols

    # If the grid is still too small to hold the layout there is nothing safe
    # to write - bail out rather than truncating the statistics silently.
    if clear_rows < layout.row_count or total_cols < NUM_COLUMNS:
        return False

    requests = _layout_to_requests(worksheet.id, layout, total_cols, clear_rows)

    spreadsheet = worksheet.spreadsheet
    for start in range(0, len(requests), REQUEST_CHUNK_SIZE):
        chunk = requests[start:start + REQUEST_CHUNK_SIZE]
        if retry:
            retry(lambda c=chunk: spreadsheet.batch_update({"requests": c}))
        else:
            spreadsheet.batch_update({"requests": chunk})

    return True


# === PRINTABLE EXCEL RENDERING ===
# The Google Sheets API has no page-setup support - no orientation, margins,
# fit-to-width or page breaks - so a Drive export can never be print-ready.
# The .xlsx is therefore built here instead, from the same layout.
def _hex(color):
    """Convert a Sheets API colour dict to an openpyxl 'FFRRGGBB' string."""
    return "FF%02X%02X%02X" % (
        round(color["red"] * 255),
        round(color["green"] * 255),
        round(color["blue"] * 255),
    )


def write_statistics_xlsx(path, sections, title, subtitle,
                          page_break_per_section=False):
    """
    Write the statistics to a print-ready .xlsx file.

    Set up so that File > Print needs no adjustment: landscape, scaled to one
    page wide, the banner repeated on every page and page numbers in the
    footer.

    Args:
        page_break_per_section: start each section (grade) on a fresh page.
            Off by default - with sixteen grades, most of which are far
            shorter than a page, that wastes most of the paper. Sections flow
            continuously instead, each still introduced by its own dark band
            and column headers.

    Returns:
        The path written, or None when there is nothing to write.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.pagebreak import Break
    from openpyxl.worksheet.properties import PageSetupProperties

    layout = build_statistics_layout(sections, title, subtitle)
    if not layout.header_rows:
        return None

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = STATS_SHEET_TITLE

    fills = {(r, c): color for r, c, color in layout.fills}
    header_rows = set(layout.header_rows)
    section_rows = set(layout.section_rows)
    subblock_rows = set(layout.subblock_rows)
    summary_rows = set(layout.summary_rows)
    title_rows = set(layout.title_rows)

    thin = Side(style="thin", color="FFD9D9D9")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    centered = Alignment(horizontal="center", vertical="center")
    left = Alignment(horizontal="left", vertical="center")

    for row_index, row in enumerate(layout.values):
        for col_index, (value, is_percent) in enumerate(row):
            cell = sheet.cell(row=row_index + 1, column=col_index + 1)

            if value is None or value == "":
                cell.value = None
            elif is_percent:
                cell.value = round(value / 100.0, 4)
                cell.number_format = "0%"
            else:
                cell.value = value

            cell.alignment = left if col_index == 0 else centered

            # Row-level banding first, then the per-cell colour coding, so a
            # summary row's own green/red cells still win.
            if row_index in title_rows:
                cell.fill = PatternFill("solid", fgColor=_hex(TITLE_BG))
                cell.font = Font(bold=True, color="FFFFFFFF", size=14)
            elif row_index in section_rows:
                cell.fill = PatternFill("solid", fgColor=_hex(SECTION_BG))
                cell.font = Font(bold=True, color="FFFFFFFF", size=12)
            elif row_index in subblock_rows:
                cell.fill = PatternFill("solid", fgColor=_hex(GREY))
                cell.font = Font(bold=True)
            elif row_index in header_rows:
                cell.fill = PatternFill("solid", fgColor=_hex(HEADER_BG))
                cell.font = Font(bold=True, color="FFFFFFFF", size=9)
                cell.border = border
                # Wrapping is what allows the narrow print widths above -
                # "Unaccounted %" becomes two lines instead of one wide column.
                cell.alignment = Alignment(
                    horizontal="center", vertical="center", wrap_text=True
                )
            elif row_index in summary_rows:
                cell.fill = PatternFill("solid", fgColor=_hex(SUMMARY_BG))
                cell.font = Font(bold=True)
                cell.border = border

            color = fills.get((row_index, col_index))
            if color is not None:
                cell.fill = PatternFill("solid", fgColor=_hex(color))
                cell.border = border
                if row_index in summary_rows:
                    cell.font = Font(bold=True)

            if col_index == 0 and row_index not in title_rows:
                cell.font = Font(
                    bold=True,
                    color="FFFFFFFF" if row_index in section_rows else "FF000000",
                    size=12 if row_index in section_rows else 11,
                )

    # Print widths, not the on-screen pixel widths - see PRINT_WIDTHS.
    for col_index, column in enumerate(COLUMNS):
        sheet.column_dimensions[get_column_letter(col_index + 1)].width = \
            PRINT_WIDTHS.get(column[1], 10)

    # Give the wrapped two-line headers room.
    for row_index in header_rows:
        sheet.row_dimensions[row_index + 1].height = 28

    # Spacer rows only need to separate grades visually, not occupy a full
    # line each - with sixteen of them that adds up on a condensed printout.
    for row_index in layout.spacer_rows:
        sheet.row_dimensions[row_index + 1].height = 6

    # The legend is a long line of prose - let it wrap rather than run off
    # the printed page.
    sheet.row_dimensions[3].height = 30
    sheet.cell(row=3, column=1).alignment = Alignment(
        horizontal="left", vertical="center", wrap_text=True
    )
    sheet.merge_cells(start_row=3, start_column=1,
                      end_row=3, end_column=NUM_COLUMNS)
    for row_index in sorted(title_rows):
        sheet.merge_cells(start_row=row_index + 1, start_column=1,
                          end_row=row_index + 1, end_column=NUM_COLUMNS)
    sheet.row_dimensions[1].height = 24

    # Optional: start each grade on a fresh page. Skips the first so the
    # banner and the opening section stay together.
    if page_break_per_section:
        for row_index in sorted(section_rows)[1:]:
            sheet.row_breaks.append(Break(id=row_index))

    # === PAGE SETUP ===
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_LETTER
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0          # as many pages tall as needed
    sheet.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    sheet.print_options.horizontalCentered = True
    sheet.page_margins.left = 0.3
    sheet.page_margins.right = 0.3
    sheet.page_margins.top = 0.5
    sheet.page_margins.bottom = 0.5
    sheet.page_margins.header = 0.2
    sheet.page_margins.footer = 0.2
    sheet.print_area = "A1:%s%d" % (
        get_column_letter(NUM_COLUMNS), layout.row_count
    )
    # Repeat the banner on every page so a loose sheet is identifiable.
    sheet.print_title_rows = "1:2"
    sheet.oddFooter.left.text = subtitle
    sheet.oddFooter.left.size = 8
    sheet.oddFooter.right.text = "Page &P of &N"
    sheet.oddFooter.right.size = 8

    sheet.freeze_panes = "B4"

    workbook.save(path)
    return path


def build_subtitle(school_year=None):
    """Standard second line: school year plus when the sheet was refreshed."""
    generated = datetime.datetime.now().strftime("%B %d, %Y at %I:%M %p")
    if school_year:
        return "School year %s   •   Last updated %s" % (school_year, generated)
    return "Last updated %s" % generated
