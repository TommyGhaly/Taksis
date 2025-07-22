import datetime
import json
import os
import tempfile
import requests
from msal import PublicClientApplication
import firebase_admin
from firebase_admin import credentials as firebase_credentials, firestore
from openpyxl import Workbook
from openpyxl.styles import PatternFill

# === CONFIG ===
with open("config.json") as f:
    config = json.load(f)

CLIENT_ID = config["client_id"]
AUTHORITY = config["authority"]
SCOPES = config["scopes"]

# === FIREBASE SETUP ===
firebase_cred = firebase_credentials.Certificate("firebase_admin.json")
firebase_admin.initialize_app(firebase_cred)
db = firestore.client()

# === ONEDRIVE LOGIN ===
app = PublicClientApplication(
    client_id=CLIENT_ID, 
    authority=AUTHORITY
)

# Try to get cached token first
accounts = app.get_accounts()
result = None

if accounts:
    print("Found cached account, attempting silent authentication...")
    try:
        result = app.acquire_token_silent(SCOPES, account=accounts[0])
        if result:
            print("✅ Using cached token")
    except Exception:
        pass

if not result:
    # Try device code flow instead of interactive
    print("Starting device code authentication...")
    try:
        flow = app.initiate_device_flow(scopes=SCOPES)
        if "user_code" not in flow:
            raise ValueError("Failed to create device flow. Check your app configuration.")
        
        print("\n" + "="*60)
        print(flow["message"])
        print("="*60 + "\n")
        
        # This will block until user completes the flow
        result = app.acquire_token_by_device_flow(flow)
    except Exception as e:
        print(f"Device flow failed: {e}")
        print("\nTrying interactive authentication as fallback...")
        try:
            result = app.acquire_token_interactive(scopes=SCOPES)
        except Exception as e2:
            print(f"Interactive auth also failed: {e2}")
            raise

# Debug: Print what we got
print(f"Authentication result keys: {result.keys() if result else 'None'}")

# Check for errors first
if result and "error" in result:
    print(f"Authentication error: {result['error']}")
    print(f"Error description: {result.get('error_description', 'No description')}")
    raise Exception(f"Authentication failed: {result['error']}")

# Get the token
access_token = result.get("access_token") if result else None
if not access_token:
    print(f"Full result object: {result}")  # Debug print
    raise Exception("Failed to retrieve access token from authentication result")

print("✅ Successfully authenticated!")

headers = {
    "Authorization": f"Bearer {access_token}",
    "Content-Type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
}

# === CONSTANTS ===
GRADES = [
    'Pre-School - Age 3', 'Pre-School - Age 4', "Kindergarten", "1st Grade", "2nd Grade",
    "3rd Grade", "4th Grade", "5th Grade", "6th Grade", "7th Grade", "8th Grade",
    "9th Grade", "10th Grade", "11th Grade", "12th Grade", "College"
]

GREEN_FILL = PatternFill(start_color="00FF00", end_color="00FF00", fill_type="solid")
RED_FILL = PatternFill(start_color="FF0000", end_color="FF0000", fill_type="solid")

# === FUNCTIONS ===
def interpret_dates(student_data):
    """Extract date entries from student data."""
    dates = []
    for key, value in student_data.items():
        try:
            datetime.datetime.strptime(key, "%m-%d-%Y")
            dates.append((key, value))
        except ValueError:
            continue
    return dates

def get_student_names(ref, grade):
    """Get all students and their attendance for a specific grade."""
    student_dict = {}
    try:
        students = ref.document(grade).collection("students").get()
        for doc in students:
            data = doc.to_dict()
            if "name" in data:
                student_dict[data["name"]] = interpret_dates(data)
    except Exception as e:
        print(f"Error getting students for {grade}: {e}")
    return student_dict

def organize_names(student_dict):
    """Sort student names alphabetically."""
    names = sorted(student_dict.keys())
    return names

def get_start_date(student_dict):
    """Find the earliest date in all attendance records."""
    date_format = "%m-%d-%Y"
    start = datetime.datetime.max
    for records in student_dict.values():
        for date_str, _ in records:
            try:
                date_obj = datetime.datetime.strptime(date_str, date_format)
                if date_obj < start:
                    start = date_obj
            except ValueError:
                continue
    return start

def create_date_array(start_date):
    """Create array of Sunday dates from start to last Sunday."""
    # Ensure start_date is a Sunday
    if start_date.weekday() != 6:
        # Adjust to previous Sunday
        days_until_sunday = (start_date.weekday() + 1) % 7
        start_date = start_date - datetime.timedelta(days=days_until_sunday)
        print(f"Adjusted start date to Sunday: {start_date.strftime('%m-%d-%Y')}")

    date_array = []
    current = start_date
    today = datetime.datetime.now()
    last_sunday = today - datetime.timedelta(days=(today.weekday() + 1) % 7)
    
    while current <= last_sunday:
        date_array.append(current.strftime("%m-%d-%Y"))
        current += datetime.timedelta(days=7)
    
    return date_array

def write_attendance_sheet(ws, student_dict, date_array, names):
    """Write attendance data to worksheet."""
    # Header row with dates
    ws.append([""] + date_array)
    
    # Student rows
    for row_idx, name in enumerate(names, start=2):
        ws.cell(row=row_idx, column=1, value=name)
        attendance = dict(student_dict[name])
        
        for col_idx, date in enumerate(date_array, start=2):
            value = attendance.get(date)
            cell = ws.cell(row=row_idx, column=col_idx)
            
            if value is True:
                cell.value = "Present"
                cell.fill = GREEN_FILL
            elif value is False:
                cell.value = "Absent"
                cell.fill = RED_FILL
            # If no value, leave cell empty

def export_church(church_doc):
    """Export attendance data for a single church."""
    data = church_doc.to_dict()
    church_name = data.get('name', 'Church')
    
    # Check if church has Sunday School sheets enabled
    if "sundaySchoolSheets" not in data or not data["sundaySchoolSheets"]:
        print(f"⏭️  Skipping {church_name} (Sunday School sheets not enabled)")
        return

    print(f"📊 Processing {church_name}...")
    
    sunday_ref = church_doc.reference.collection("Sunday School")
    wb = Workbook()
    wb.remove(wb.active)  # Remove default sheet
    
    sheets_created = 0
    
    for grade in GRADES:
        student_dict = get_student_names(sunday_ref, grade)
        if not student_dict:
            continue
        
        print(f"  - Adding {grade} ({len(student_dict)} students)")
        
        sheet = wb.create_sheet(title=grade)
        names = organize_names(student_dict)
        
        try:
            start_date = get_start_date(student_dict)
            date_array = create_date_array(start_date)
            write_attendance_sheet(sheet, student_dict, date_array, names)
            sheets_created += 1
        except Exception as e:
            print(f"    ⚠️  Error processing {grade}: {e}")

    if sheets_created == 0:
        print(f"  ⚠️  No attendance data found for {church_name}")
        return

    # Save temp file and upload
    with tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx") as tmp:
        wb.save(tmp.name)
        tmp.close()

        upload_name = f"{church_name}_Attendance.xlsx"
        try:
            with open(tmp.name, "rb") as f:
                res = requests.put(
                    f"https://graph.microsoft.com/v1.0/me/drive/root:/Documents/{upload_name}:/content",
                    headers=headers,
                    data=f.read()
                )
                
                if res.status_code in [200, 201]:
                    print(f"  ✅ Uploaded {upload_name}")
                else:
                    print(f"  ❌ Failed to upload {upload_name}: {res.status_code}")
                    if res.text:
                        error_data = res.json()
                        print(f"     Error: {error_data.get('error', {}).get('message', 'Unknown error')}")
        except Exception as e:
            print(f"  ❌ Upload error: {e}")
        finally:
            os.remove(tmp.name)

# === MAIN EXECUTION ===
def main():
    print("🚀 Starting attendance sheet export...")
    print(f"📅 Processing attendance up to: {datetime.datetime.now().strftime('%m-%d-%Y')}")
    print()
    
    try:
        churches = db.collection("churches").get()
        print(f"Found {len(churches)} churches in database")
        print()
        
        for church in churches:
            try:
                export_church(church)
            except Exception as e:
                church_name = church.to_dict().get('name', 'Unknown')
                print(f"❌ Error processing {church_name}: {e}")
            print()  # Blank line between churches
        
        print("✅ Export complete!")
        
    except Exception as e:
        print(f"❌ Fatal error: {e}")
        raise

if __name__ == "__main__":
    main()